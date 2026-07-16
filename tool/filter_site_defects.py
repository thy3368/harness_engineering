from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl


SOURCE = Path("task/列表_现场缺陷1-6月_2026-07-09 15-27-47.xlsx")
OUTPUT = Path("outputs/filtered_现场缺陷_指定编号_2026-07-10.xlsx")
SHEET_NAME = "需求列表"

TARGET_IDS = [
    "202602023427",
    "202602034483",
    "202602104677",
    "202602263472",
    "202603053030",
    "202603113790",
    "202603114586",
    "202603194230",
    "202603194251",
    "202603264635",
    "202604303270",
    "202605093047",
    "202605133768",
    "202605193245",
    "202606080820",
    "202606120664",
    "202606180485",
    "202606181100",
    "202606250062",
]


def copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
    if src.hyperlink:
        dst.hyperlink = src.hyperlink.target
        if src.hyperlink.tooltip:
            dst.hyperlink.tooltip = src.hyperlink.tooltip
    if src.comment:
        dst.comment = copy(src.comment)


def main() -> None:
    wb = openpyxl.load_workbook(SOURCE)
    ws = wb[SHEET_NAME]

    target_set = set(TARGET_IDS)
    rows_by_id: dict[str, int] = {}
    matched_source_rows: list[int] = []
    duplicates: list[str] = []

    for row_idx in range(3, ws.max_row + 1):
        value = ws.cell(row_idx, 1).value
        defect_id = str(value).strip() if value is not None else ""
        if defect_id in target_set:
            if defect_id in rows_by_id:
                duplicates.append(defect_id)
            rows_by_id[defect_id] = row_idx
            matched_source_rows.append(row_idx)

    missing = [defect_id for defect_id in TARGET_IDS if defect_id not in rows_by_id]
    if missing:
        raise SystemExit(f"Missing target ids in source workbook: {missing}")
    if duplicates:
        raise SystemExit(f"Duplicate target ids in source workbook: {duplicates}")

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = SHEET_NAME

    source_rows = [1, 2, *matched_source_rows]
    for out_row_idx, source_row_idx in enumerate(source_rows, start=1):
        for col_idx in range(1, 14):
            copy_cell(ws.cell(source_row_idx, col_idx), out_ws.cell(out_row_idx, col_idx))

        source_dim = ws.row_dimensions[source_row_idx]
        target_dim = out_ws.row_dimensions[out_row_idx]
        target_dim.height = source_dim.height
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed

    for col_idx in range(1, 14):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        source_dim = ws.column_dimensions[col_letter]
        target_dim = out_ws.column_dimensions[col_letter]
        target_dim.width = source_dim.width
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel
        target_dim.collapsed = source_dim.collapsed

    if ws.freeze_panes:
        out_ws.freeze_panes = ws.freeze_panes

    out_ws.auto_filter.ref = f"A2:M{out_ws.max_row}"
    out_ws.sheet_view.showGridLines = ws.sheet_view.showGridLines
    out_ws.sheet_format.defaultRowHeight = ws.sheet_format.defaultRowHeight
    out_ws.sheet_format.defaultColWidth = ws.sheet_format.defaultColWidth

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(OUTPUT)

    check_wb = openpyxl.load_workbook(OUTPUT, read_only=True, data_only=True)
    check_ws = check_wb[SHEET_NAME]
    output_ids = [
        str(check_ws.cell(row_idx, 1).value).strip()
        for row_idx in range(3, check_ws.max_row + 1)
    ]

    if check_ws.max_row != 21:
        raise SystemExit(f"Expected 21 rows, got {check_ws.max_row}")
    if output_ids != [str(ws.cell(row_idx, 1).value).strip() for row_idx in matched_source_rows]:
        raise SystemExit("Output row order does not match source row order")
    if set(output_ids) != target_set:
        raise SystemExit("Output ids do not match target id set")
    if len(output_ids) != len(set(output_ids)):
        raise SystemExit("Output contains duplicate ids")

    print(f"saved={OUTPUT}")
    print(f"rows={check_ws.max_row}")
    print(f"data_rows={len(output_ids)}")
    print("ids=" + ",".join(output_ids))


if __name__ == "__main__":
    main()
