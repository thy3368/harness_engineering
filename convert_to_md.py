#!/usr/bin/env python3
import os
import sys
import subprocess
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    '.pdf': 'PDF文档',
    '.docx': 'Word文档',
    '.xlsx': 'Excel表格',
}

def check_install_deps():
    deps_ok = True
    try:
        import docx
    except ImportError:
        print("安装 python-docx...")
        subprocess.run([sys.executable, '-m', 'pip', 'install', 'python-docx'], capture_output=True)
        import docx
    
    try:
        import openpyxl
    except ImportError:
        print("安装 openpyxl...")
        subprocess.run([sys.executable, '-m', 'pip', 'install', 'openpyxl'], capture_output=True)
        import openpyxl
    
    try:
        import pymupdf
    except ImportError:
        print("安装 pymupdf...")
        subprocess.run([sys.executable, '-m', 'pip', 'install', 'pymupdf'], capture_output=True)
        import pymupdf

def convert_pdf(file_path: Path) -> str:
    import pymupdf
    doc = pymupdf.open(str(file_path))
    text = ""
    for page in doc:
        text += page.get_text() + "\n"
    doc.close()
    return text

def convert_docx(file_path: Path) -> str:
    from docx import Document
    doc = Document(str(file_path))
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def convert_xlsx(file_path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows():
            row_text = []
            for cell in row:
                val = cell.value
                if val is not None:
                    row_text.append(str(val))
            if row_text:
                text += " | ".join(row_text) + "\n"
        text += "\n"
    wb.close()
    return text

def get_files_to_convert(directory: str) -> list:
    files = []
    dir_path = Path(directory)
    
    if not dir_path.exists():
        print(f"目录不存在: {directory}")
        return files
    
    for ext in SUPPORTED_EXTENSIONS:
        files.extend(dir_path.rglob(f"*{ext}"))
    
    md_files = set(dir_path.rglob("*.md"))
    files = [f for f in files if f not in md_files]
    
    return sorted(files)

def convert_file(file_path: Path) -> bool:
    output_path = file_path.parent / f"{file_path.stem}.md"
    
    if output_path.exists():
        print(f"  ⏭️  跳过 (已存在): {output_path.name}")
        return True
    
    print(f"  🔄 转换: {file_path.name}")
    
    try:
        ext = file_path.suffix.lower()
        if ext == '.pdf':
            content = convert_pdf(file_path)
        elif ext == '.docx':
            content = convert_docx(file_path)
        elif ext == '.xlsx':
            content = convert_xlsx(file_path)
        else:
            print(f"  ❌ 不支持的格式: {ext}")
            return False
        
        output_path.write_text(content, encoding='utf-8')
        print(f"  ✅ 完成: {output_path.name}")
        return True
    except Exception as e:
        print(f"  ❌ 失败: {file_path.name} - {str(e)}")
        return False

def main():
    target_dir = "/Users/hongyaotang/src/harness_engineering/现有规范/基础层规范"
    
    print(f"📂 目标目录: {target_dir}")
    print("=" * 60)
    
    print("🔧 检查依赖...")
    check_install_deps()
    
    files = get_files_to_convert(target_dir)
    
    if not files:
        print("✅ 没有找到需要转换的文件")
        return
    
    print(f"📋 找到 {len(files)} 个文件需要转换:\n")
    for f in files:
        ext = f.suffix.lower()
        file_type = SUPPORTED_EXTENSIONS.get(ext, '未知')
        print(f"  - {f.name} ({file_type})")
    
    print("\n" + "=" * 60)
    print("🚀 开始转换...\n")
    
    success = 0
    failed = 0
    
    for file_path in files:
        if convert_file(file_path):
            success += 1
        else:
            failed += 1
    
    print("\n" + "=" * 60)
    print(f"📊 转换完成: 成功 {success}, 失败 {failed}")

if __name__ == "__main__":
    main()
